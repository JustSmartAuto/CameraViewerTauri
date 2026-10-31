"""SQLite + xlsx 双路落库。

SQLite 增量写入所有采样点（时间戳、相机、单元格、值、OK/NG）；
xlsx 按 相机/日期/单元格 导出工作表，文件名带时间戳后缀。
"""

from __future__ import annotations

import sqlite3
import threading
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

SCHEMA = """
CREATE TABLE IF NOT EXISTS samples (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    ts TEXT NOT NULL,               -- ISO8601 本地时间
    camera TEXT NOT NULL,           -- 相机名
    cell TEXT NOT NULL,             -- 单元格名/地址
    value REAL,                     -- 数值（无法解析为数值时为 NULL）
    text TEXT,                      -- 原始文本
    ok INTEGER                      -- 1=合格 0=超差 NULL=无阈值
);
CREATE INDEX IF NOT EXISTS idx_samples_cam_cell ON samples(camera, cell, ts);

CREATE TABLE IF NOT EXISTS images (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    ts TEXT NOT NULL,               -- ISO8601 本地时间（同一次拍照共用）
    camera TEXT NOT NULL,
    cell TEXT NOT NULL,
    kind TEXT NOT NULL,             -- image=原始图 render=结果渲染图
    mime TEXT,                      -- 如 image/png
    width INTEGER,
    height INTEGER,
    data BLOB NOT NULL              -- 图像二进制
);
CREATE INDEX IF NOT EXISTS idx_images_cam_cell ON images(camera, cell, ts);
"""


def timestamp_suffix(ts: datetime | None = None) -> str:
    return (ts or datetime.now()).strftime("_%Y%m%d_%H%M%S")


class Storage:
    """SQLite 写入器（线程安全，供 asyncio 线程调用）。"""

    def __init__(self, db_path: Path):
        db_path = Path(db_path)
        db_path.parent.mkdir(parents=True, exist_ok=True)
        self._lock = threading.Lock()
        self._conn = sqlite3.connect(str(db_path), check_same_thread=False)
        self._conn.executescript(SCHEMA)
        self._conn.commit()

    def insert_sample(
        self,
        ts: datetime,
        camera: str,
        cell: str,
        raw,
        ok: bool | None,
    ) -> None:
        value = None
        try:
            value = float(raw)
        except (TypeError, ValueError):
            pass
        with self._lock:
            self._conn.execute(
                "INSERT INTO samples (ts, camera, cell, value, text, ok) VALUES (?,?,?,?,?,?)",
                (ts.isoformat(timespec="milliseconds"), camera, cell, value, str(raw), ok),
            )
            self._conn.commit()

    def insert_image(
        self,
        ts: datetime,
        camera: str,
        cell: str,
        kind: str,
        data: bytes,
        mime: str = "image/png",
    ) -> tuple[int | None, int | None]:
        """写入图像 BLOB，返回 (宽, 高)（无法解析时为 None）。"""
        width = height = None
        try:
            from io import BytesIO

            from PIL import Image as PILImage

            with PILImage.open(BytesIO(data)) as im:
                width, height = im.size
        except Exception:
            pass
        with self._lock:
            self._conn.execute(
                "INSERT INTO images (ts, camera, cell, kind, mime, width, height, data)"
                " VALUES (?,?,?,?,?,?,?,?)",
                (ts.isoformat(timespec="milliseconds"), camera, cell, kind,
                 mime, width, height, data),
            )
            self._conn.commit()
        return width, height

    def query_images(
        self,
        camera: str,
        cell: str,
        kind: str | None = None,
        start: datetime | None = None,
        end: datetime | None = None,
    ) -> list[tuple[str, str, str, int | None, int | None, bytes]]:
        """返回 [(ts, kind, mime, width, height, data)]，按时间升序。"""
        sql = "SELECT ts, kind, mime, width, height, data FROM images WHERE camera=? AND cell=?"
        args: list = [camera, cell]
        if kind:
            sql += " AND kind=?"
            args.append(kind)
        if start:
            sql += " AND ts>=?"
            args.append(start.isoformat(timespec="milliseconds"))
        if end:
            sql += " AND ts<=?"
            args.append(end.isoformat(timespec="milliseconds"))
        sql += " ORDER BY ts"
        with self._lock:
            return self._conn.execute(sql, args).fetchall()

    def query(
        self,
        camera: str,
        cell: str,
        start: datetime | None = None,
        end: datetime | None = None,
    ) -> list[tuple[str, float | None, str, int | None]]:
        """返回 [(ts, value, text, ok)]，按时间升序。"""
        sql = "SELECT ts, value, text, ok FROM samples WHERE camera=? AND cell=?"
        args: list = [camera, cell]
        if start:
            sql += " AND ts>=?"
            args.append(start.isoformat(timespec="milliseconds"))
        if end:
            sql += " AND ts<=?"
            args.append(end.isoformat(timespec="milliseconds"))
        sql += " ORDER BY ts"
        with self._lock:
            return self._conn.execute(sql, args).fetchall()

    def cells(self, camera: str) -> list[str]:
        with self._lock:
            rows = self._conn.execute(
                "SELECT DISTINCT cell FROM samples WHERE camera=? ORDER BY cell", (camera,)
            ).fetchall()
        return [r[0] for r in rows]

    def query_all(
        self,
        camera: str,
        start: datetime | None = None,
        end: datetime | None = None,
    ) -> list[tuple[str, str, float | None, str, int | None]]:
        """返回某相机全部采样点 [(ts, cell, value, text, ok)]，按时间升序。"""
        sql = "SELECT ts, cell, value, text, ok FROM samples WHERE camera=?"
        args: list = [camera]
        if start:
            sql += " AND ts>=?"
            args.append(start.isoformat(timespec="milliseconds"))
        if end:
            sql += " AND ts<=?"
            args.append(end.isoformat(timespec="milliseconds"))
        sql += " ORDER BY ts, cell"
        with self._lock:
            return self._conn.execute(sql, args).fetchall()

    def cameras(self) -> list[str]:
        with self._lock:
            rows = self._conn.execute("SELECT DISTINCT camera FROM samples ORDER BY camera").fetchall()
        return [r[0] for r in rows]

    def close(self) -> None:
        with self._lock:
            self._conn.close()


def export_xlsx(
    storage: Storage,
    camera: str,
    out_dir: Path,
    start: datetime | None = None,
    end: datetime | None = None,
    notes: dict[str, str] | None = None,
    image_cells: set[str] | None = None,
) -> Path:
    """按 单元格 生成工作表导出（文件名带时间戳后缀）。

    图像单元格：直接嵌入图像（原始图工作表 + A0 的渲染图工作表）。
    额外生成「合表」工作表：将同一时间戳（同一次拍照）下所有单元格的值
    合并到同一行，列名为单元格名，便于观察同一次检测中各数据的相互关系。
    """
    from collections import defaultdict

    wb = Workbook()
    wb.remove(wb.active)
    red = PatternFill("solid", start_color="FFC7CE")
    bold = Font(bold=True)
    notes = notes or {}
    image_cells = set(image_cells or ())
    cell_names = storage.cells(camera)
    for cell in cell_names:
        if cell in image_cells:
            _append_image_sheets(wb, storage, camera, cell, notes.get(cell, ""), bold,
                                 start, end)
            continue
        rows = storage.query(camera, cell, start, end)
        ws = wb.create_sheet(title=cell[:31])
        ws.append(["时间戳", "数值", "原始值", "判定"])
        for c in ws[1]:
            c.font = bold
        if notes.get(cell):
            ws["F1"] = "注解"
            ws["F1"].font = bold
            ws["F2"] = notes[cell]
        for ts, value, text, ok in rows:
            ws.append([ts, value, text, "OK" if ok else ("NG" if ok == 0 else "")])
            if ok == 0:
                for c in ws[ws.max_row]:
                    c.fill = red
        if ws.max_row > 1:
            ws.auto_filter.ref = f"A1:D{ws.max_row}"

    # 合表：同一次拍照（同一时间戳）的所有单元格合并到一行
    if cell_names:
        all_rows = storage.query_all(camera, start, end)
        grouped: dict[str, dict[str, tuple[float | None, str, int | None]]] = defaultdict(dict)
        for ts, cell, value, text, ok in all_rows:
            grouped[ts][cell] = (value, text, ok)
        ws = wb.create_sheet(title="合表", index=0)
        headers = ["时间戳"]
        for cell in cell_names:
            note = notes.get(cell, "")
            headers.append(f"{cell}（{note}）" if note else cell)
        ws.append(headers)
        for c in ws[1]:
            c.font = bold
        for ts in sorted(grouped):
            row_vals = grouped[ts]
            # 优先显示数值，无法转数值时显示原始文本
            row = [ts]
            for cell in cell_names:
                item = row_vals.get(cell)
                if item is None:
                    row.append(None)
                else:
                    value, text, _ok = item
                    row.append(value if value is not None else text)
            ws.append(row)
        if ws.max_row > 1:
            ws.auto_filter.ref = f"A1:{ws.cell(row=1, column=len(cell_names)+1).coordinate}"
            ws.freeze_panes = "B2"

    out_dir.mkdir(parents=True, exist_ok=True)
    path = out_dir / f"{camera}{timestamp_suffix()}.xlsx"
    wb.save(path)
    return path


def _embed_image(ws, row: int, data: bytes, max_width: int = 480) -> bool:
    """把图像 BLOB 嵌入单元格（B{row}），返回是否成功。"""
    from io import BytesIO

    from openpyxl.drawing.image import Image as XLImage

    try:
        img = XLImage(BytesIO(data))
    except Exception:
        return False
    if img.width > max_width:
        ratio = max_width / img.width
        img.width = max_width
        img.height = int(img.height * ratio)
    ws.add_image(img, f"B{row}")
    # 行高按像素换算为磅（1px ≈ 0.75pt），避免图像压到下一行
    ws.row_dimensions[row].height = max(15, img.height * 0.75)
    return True


def _append_image_sheets(wb, storage: Storage, camera: str, cell: str, note: str,
                         bold, start, end) -> None:
    """图像单元格：原始图与渲染图各生成一个嵌入式工作表。"""
    rows = storage.query_images(camera, cell, start=start, end=end)
    for kind, title in (("image", cell[:31]), ("render", f"{cell}渲染图"[:31])):
        items = [r for r in rows if r[1] == kind]
        if not items:
            continue
        ws = wb.create_sheet(title=title)
        ws.column_dimensions["A"].width = 26
        ws.column_dimensions["B"].width = 72
        ws.append(["时间戳", "图像"])
        for c in ws[1]:
            c.font = bold
        if note:
            ws["D1"] = "注解"
            ws["D1"].font = bold
            ws["D2"] = note
        for i, (ts, _kind, _mime, width, height, data) in enumerate(items, start=2):
            ws.cell(row=i, column=1, value=ts)
            if not _embed_image(ws, i, data):
                ws.cell(row=i, column=2, value=f"[图像 {width}×{height}]")
        ws.freeze_panes = "A2"
